from pathlib import Path
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm


pdf = Path("策划书-弈金-数字金融版.pdf")
out = Path("策划书-弈金-数字金融版.docx")
image_dir = Path("/tmp/yijin_pdf_pages")
image_dir.mkdir(parents=True, exist_ok=True)

subprocess.run(
    ["pdftoppm", "-png", "-r", "160", str(pdf), str(image_dir / "page")],
    check=True,
)

doc = Document()
for section in doc.sections:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0)
    section.bottom_margin = Cm(0)
    section.left_margin = Cm(0)
    section.right_margin = Cm(0)

pages = sorted(image_dir.glob("page-*.png"))
for index, page in enumerate(pages):
    if index:
        doc.add_section(WD_SECTION_START.NEW_PAGE)
        section = doc.sections[-1]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(0)
        section.bottom_margin = Cm(0)
        section.left_margin = Cm(0)
        section.right_margin = Cm(0)
    paragraph = doc.paragraphs[0] if index == 0 else doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Cm(0)
    paragraph.paragraph_format.space_after = Cm(0)
    paragraph.add_run().add_picture(str(page), width=Cm(21), height=Cm(29.7))

doc.core_properties.title = "弈金 数字金融项目策划书"
doc.core_properties.subject = "PDF 高保真 Word 版本"
doc.save(out)
print(out)
