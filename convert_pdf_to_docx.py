from pathlib import Path
import re
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


PDF = Path("策划书-弈金-数字金融版.pdf")
DOCX = Path("策划书-弈金-数字金融版.docx")


def set_run_font(run, east_asia="宋体", size=10.5, bold=False):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def is_heading(text):
    return bool(
        re.match(r"^(项目摘要|[一二三四五六七八九十]+、|第十七届|弈金：|附录|[0-9]+\.[0-9]+\s+)", text)
    )


def clean_page_lines(raw, page_index):
    lines = [re.sub(r"[ \t]+$", "", x) for x in raw.splitlines()]
    lines = [x for x in lines if x.strip()]
    if page_index > 0 and lines and lines[0].startswith("弈金·工行杯策划书"):
        lines.pop(0)
    if lines and re.fullmatch(r"—[0-9]+—", lines[-1].strip()):
        lines.pop()
    return lines


def add_paragraph(doc, text, style="Body Text", align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    safe_text = "".join(ch for ch in text.strip() if ch in "\n\t" or ord(ch) >= 32)
    r = p.add_run(safe_text)
    set_run_font(r, size=11 if style in ("Title", "Heading 1") else 10.5,
                 bold=style in ("Title", "Heading 1", "Heading 2"))
    return p


extracted = subprocess.run(
    ["pdftotext", "-layout", str(PDF), "-"],
    check=True,
    capture_output=True,
    text=True,
).stdout
page_texts = extracted.split("\f")
doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
styles["Normal"].font.size = Pt(10.5)
for name, size, space_before, space_after in [
    ("Title", 24, 0, 14),
    ("Heading 1", 16, 14, 8),
    ("Heading 2", 13, 10, 5),
]:
    st = styles[name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = None
    st.paragraph_format.space_before = Pt(space_before)
    st.paragraph_format.space_after = Pt(space_after)

body = styles["Body Text"]
body.font.name = "Aptos"
body._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
body.font.size = Pt(10.5)
body.paragraph_format.first_line_indent = Cm(0.74)
body.paragraph_format.line_spacing = 1.25
body.paragraph_format.space_after = Pt(4)

bullet = styles["List Bullet"]
bullet.font.name = "Aptos"
bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
bullet.font.size = Pt(10.5)
bullet.paragraph_format.left_indent = Cm(0.74)
bullet.paragraph_format.space_after = Pt(2)

first_page = True
for page_index, raw in enumerate(page_texts):
    lines = clean_page_lines(raw, page_index)
    if page_index:
        doc.add_page_break()

    # Keep the cover page centered and visually distinct.
    if first_page:
        first_page = False
        for line in lines:
            t = line.strip()
            if not t:
                continue
            if t in ("弈金", "数字金融项目策划书"):
                p = add_paragraph(doc, t, "Title", WD_ALIGN_PARAGRAPH.CENTER)
                p.paragraph_format.space_before = Pt(18 if t == "弈金" else 4)
            elif "面向银行投研场景" in t or "核心命题" in t:
                add_paragraph(doc, t, "Body Text", WD_ALIGN_PARAGRAPH.CENTER)
            else:
                add_paragraph(doc, t, "Body Text", WD_ALIGN_PARAGRAPH.CENTER)
        continue

    # Merge wrapped PDF lines into paragraphs while preserving explicit tables.
    current = []
    for line in lines:
        text = line.strip()
        if not text:
            if current:
                add_paragraph(doc, "".join(current))
                current = []
            continue
        if text.startswith("·") or text.startswith("•"):
            if current:
                add_paragraph(doc, "".join(current))
                current = []
            add_paragraph(doc, text.lstrip("·• "), "List Bullet")
            continue
        if is_heading(text):
            if current:
                add_paragraph(doc, "".join(current))
                current = []
            style = "Heading 1" if re.match(r"^(项目摘要|[一二三四五六七八九十]+、|附录)", text) else "Heading 2"
            add_paragraph(doc, text, style)
            continue
        # Preserve multi-column value rows as text with tab stops rather than collapsing them.
        if len(line) - len(line.lstrip()) > 2 and ("          " in line or "     " in line):
            if current:
                add_paragraph(doc, "".join(current))
                current = []
            p = add_paragraph(doc, re.sub(r"\s{2,}", "\t", line.strip()))
            p.paragraph_format.first_line_indent = Cm(0)
            continue
        current.append(text)
    if current:
        add_paragraph(doc, "".join(current))

doc.core_properties.title = "弈金 数字金融项目策划书"
doc.core_properties.subject = "PDF 转 Word 可编辑版本"
doc.save(DOCX)
print(DOCX)
