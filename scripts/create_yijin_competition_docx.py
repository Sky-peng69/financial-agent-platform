from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('/Users/laurence/Documents/金融agent/docs/competition/弈金-工行杯项目策划书.docx')
BLUE = '173A7A'; LIGHT_BLUE = 'EAF1FA'; RED = 'C9252D'; GRAY = 'D9DDE6'; DARK = '111827'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn('w:shd'))
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color=GRAY, size='6'):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge; el = borders.find(qn(tag))
        if el is None: el = OxmlElement(tag); borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),size); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)

def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn('w:'+m))
        if node is None: node = OxmlElement('w:'+m); mar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def set_run_font(run, name='PingFang SC', size=11, bold=False, color=DARK):
    run.font.name = name; run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = RGBColor.from_string(color)

def add_text(p, text, bold=False, size=11, color=DARK):
    r = p.add_run(text); set_run_font(r, size=size, bold=bold, color=color); return r

def add_para(doc, text='', style='Body Text', first=True, space_after=6):
    p = doc.add_paragraph(style=style); p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.line_spacing = 1.35
    if first and text: p.paragraph_format.first_line_indent = Cm(0.74)
    add_text(p, text); return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.25; add_text(p,text); return p

def add_heading(doc, text, level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.keep_with_next=True; add_text(p,text,bold=(level==1),size=16 if level==1 else 13,color=DARK); return p

def add_table(doc, headers, rows):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'
    for j,h in enumerate(headers):
        c=table.rows[0].cells[j]; set_cell_shading(c,BLUE); set_cell_border(c); set_cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_text(p,h,bold=True,size=9,color='FFFFFF')
    for i,row in enumerate(rows):
        cells=table.add_row().cells
        for j,val in enumerate(row):
            c=cells[j]; set_cell_border(c); set_cell_margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i%2==1: set_cell_shading(c,LIGHT_BLUE)
            p=c.paragraphs[0]; p.paragraph_format.line_spacing=1.15; add_text(p,str(val),size=9)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return table

def page_break(doc): doc.add_page_break()

def add_source(doc, label, url):
    p=doc.add_paragraph(style='Body Text'); p.paragraph_format.space_after=Pt(3); p.paragraph_format.left_indent=Cm(.74); add_text(p,f'{label}：{url}',size=9,color='555555')

doc=Document(); sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.top_margin=Cm(2.7); sec.bottom_margin=Cm(2.4); sec.left_margin=Cm(2.8); sec.right_margin=Cm(2.8)
styles=doc.styles
for name in ['Normal','Body Text']:
    st=styles[name]; st.font.name='PingFang SC'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'PingFang SC'); st.font.size=Pt(11); st.font.color.rgb=RGBColor.from_string(DARK)
for level,size in [(1,16),(2,13),(3,11)]:
    st=styles[f'Heading {level}']; st.font.name='PingFang SC'; st._element.rPr.rFonts.set(qn('w:eastAsia'),st.font.name); st.font.size=Pt(size); st.font.bold=level==1; st.font.color.rgb=RGBColor.from_string(DARK); st.paragraph_format.space_before=Pt(12 if level==1 else 8); st.paragraph_format.space_after=Pt(5)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_text(footer,'弈金｜工行杯项目策划书  ·  ',size=9,color='777777'); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(90); add_text(p,'第十七届“工行杯”全国大学生金融科技创新大赛',bold=True,size=17,color=BLUE)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(34); add_text(p,'弈金',bold=True,size=36,color=DARK)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_text(p,'面向金融机构的可追溯研究智能体工作台',size=18,color=RED)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(32); add_text(p,'项目策划书',size=16,color=DARK)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(110); add_text(p,'参赛方向建议：数字金融 / 金融安全 / 开放创新',size=11,color='555555')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_text(p,'版本：2026 年 9 月',size=10,color='777777')
page_break(doc)

add_heading(doc,'项目摘要',1)
add_para(doc,'弈金是一款面向证券公司、基金公司及银行研究场景的金融智能体工作台。用户可以上传年报、公告等 PDF 材料，提出一个围绕 A 股上市公司的研究问题，由 Commander 将问题拆解为结构化任务，再调度行业、基本面、消息面、估值、风险和报告合成等 Specialist 并行或按依赖执行。系统保存搜索来源、页码定位、任务状态、Agent 输出、质量状态和待复核项，生成带来源和风险提示的研究报告草稿。')
add_para(doc,'项目的核心不是让模型替代研究员，而是把资料搜集、文档阅读、信息整理、初步分析和报告编排中高重复的部分组织起来，让研究人员把时间用于判断、复核和决策。产品当前定位为研究辅助工具，不自动交易，不生成面向个人用户的个性化买卖指令，不把模型输出视为未经复核的正式金融结论。')
add_table(doc,['维度','项目回答'],[
    ['参赛价值','以 AI、多智能体协作和证据链技术改造金融研究工作流，回应数智银行、金融安全和开放创新方向。'],
    ['目标用户','证券公司和基金公司的投研人员；银行研究、科技金融尽调、企业客户分析和风险复核岗位。'],
    ['产品形态','Web 工作台，支持 SaaS 与 Docker 私有化部署；前端展示任务进度、引用和报告，后台 Worker 执行长任务。'],
    ['当前边界','正式支持 PDF；公开网页检索为主要外部数据来源；结构化财务数据接口、OCR 和自动交易不纳入当前首版。'],
    ['工行适配','可作为工银智涌、多模型、知识工程和智能体协同能力的业务编排与研究应用层候选方案，具体接入需经过授权和安全评估。'],
])
page_break(doc)

add_heading(doc,'一 项目概况与参赛定位',1)
add_heading(doc,'1.1 项目名称与一句话定位',2)
add_para(doc,'项目名称：弈金。')
add_para(doc,'一句话定位：让金融研究人员从“找资料、做整理、写初稿”中解放出来，用一支可编排、可追溯、可人工复核的 AI 分析团队完成研究准备工作。')
add_heading(doc,'1.2 与本届赛事的匹配关系',2)
add_para(doc,'第十七届“工行杯”以“数智银行、创见未来”为主题，设置十大参赛方向。弈金不把自己包装为泛化的聊天机器人，而是聚焦金融机构真实工作流，建议以数字金融为主方向，以金融安全、开放创新和智慧银行为交叉价值。')
add_table(doc,['赛事关注点','弈金的对应回答','当前证据'],[
    ['真实金融场景','围绕 A 股公司研究、科技金融尽调、企业客户分析与风险复核。','已有产品定位、页面和任务编排代码'],
    ['AI 技术应用','Commander + Specialist、SSE 实时进度、可替换模型 Provider。','现有后端编排、Agent 与 LLM 服务'],
    ['金融安全','来源回链、人工复核、风险提示、权限隔离和不自动交易。','产品规范与现有引用展示'],
    ['落地可行性','SaaS + Docker 私有化；先做研究闭环，再扩展行业和风控场景。','docker-compose 与项目路线图'],
])
add_para(doc,'事实边界：本项目与工商银行公开的数智化方向进行场景对齐，不代表已经获得工商银行授权、接入工商银行内部系统、使用工商银行非公开数据或形成商业合作。后续若进入试点，必须经过数据、模型、系统和业务安全评估。', first=False)

add_heading(doc,'二 背景分析与需求判断',1)
add_heading(doc,'2.1 金融研究的工作流痛点',2)
add_bullet(doc,'资料来源分散：年报、公告、监管信息、新闻和行业资料分散在多个入口，研究人员需要反复搜索、下载、阅读和整理。')
add_bullet(doc,'多维分析难以并行：行业、基本面、消息面、估值和风险之间既有专业分工，又需要最终交叉验证。')
add_bullet(doc,'证据链不完整：普通对话式 AI 往往只给出结论，不能稳定返回页码、表格位置、来源地址、获取时间和计算过程。')
add_bullet(doc,'研究结果难以复用：一次任务的中间过程、失败原因和人工修改点没有统一沉淀，刷新或换人后容易重复劳动。')
add_bullet(doc,'金融容错率低：模型幻觉、数据口径不一致和引用失真，会直接影响研究判断，因此系统必须保留人工复核边界。')
add_heading(doc,'2.2 赛事与工商银行公开方向',2)
add_para(doc,'赛事官方信息显示，本届大赛鼓励参赛者结合真实金融业务场景，使用 AI 等新技术提出内容充实、创意新颖、可行性强的方案。工商银行 2025 年度报告披露了“领航AI+”行动、“工银智涌”大模型技术体系、多模型融合、企业知识工程、智能体协同和基础设施/数据/模型/应用全链路安全防护等方向。')
add_para(doc,'这些公开信息给弈金的启发是：应用层产品不应只强调“模型很强”，而应说明知识如何进入系统、任务如何组织、风险如何控制、结果如何复核，以及如何在现有金融科技底座之上形成可部署的业务工作流。')

add_heading(doc,'三 产品方案',1)
add_heading(doc,'3.1 产品闭环',2)
add_table(doc,['阶段','用户动作','系统动作','交付物'],[
    ['创建任务','填写公司、问题和研究范围；上传 PDF。','校验文件并建立用户、机构、任务、文件关系。','研究任务与证据集合'],
    ['任务规划','确认研究目标。','Commander 识别目标，生成带 Agent、输入、输出和依赖的任务图。','结构化 DAG'],
    ['专业执行','查看进度，必要时取消或重试。','Specialist 并行或按依赖执行；Search Provider 保存公开来源。','专业分析与引用'],
    ['统一审查','查看来源、质量状态和待复核项。','共享协议执行数字、引用、完整性和风险检查。','证据与质量状态'],
    ['报告复核','编辑、确认或退回草稿。','报告合成师输出 Markdown/PDF；未复核内容保留草稿标识。','可导出研究报告'],
])
add_heading(doc,'3.2 核心功能',2)
add_bullet(doc,'智能分析指挥台：自然语言输入问题，Commander 自动判断简单任务与复杂任务，按需选择 Specialist。')
add_bullet(doc,'多 Agent 分析引擎：行业分析、基本面分析、消息面分析、估值/财务建模、风险分析和报告合成等核心角色。')
add_bullet(doc,'PDF 解析与证据检索：支持文本提取、文档分块、关键词检索、页码定位和引用回链。')
add_bullet(doc,'实时进度与任务恢复：后台 Worker 执行长任务，通过 SSE 展示状态；用户刷新后仍可恢复查看。')
add_bullet(doc,'报告与人工复核：在线 Markdown 展示，支持 PDF 导出；报告包含来源、生成时间、风险提示和人工复核状态。')
add_bullet(doc,'模型和搜索解耦：模型通过统一 Provider 接口调用，联网检索由 Search Provider 负责，避免 Agent 绑定单一 SDK。')
add_heading(doc,'3.3 产品当前能力与后续规划',2)
add_table(doc,['能力层','当前可说明内容','下一阶段规划','边界'],[
    ['任务编排','Commander + Specialist 结构、SSE、任务历史。','结构化 DAG 依赖、重试策略和取消恢复增强。','不允许无限重试和无限 Agent 派生。'],
    ['数据来源','上传 PDF 与公开网页检索。','接入可插拔的公开结构化金融数据源。','关键财务数字不得只依赖普通搜索摘要。'],
    ['模型','DeepSeek 为当前默认 Provider。','支持更多国产或本地模型。','Agent 不直接依赖具体 SDK。'],
    ['文件','PDF 文本与引用定位。','表格抽取和更强的文档质量检测。','OCR、Word、Excel 不属于当前首版。'],
    ['报告','在线 Markdown、PDF 导出、人工复核标记。','图表、评测集和导出一致性增强。','Word 导出不属于当前首版。'],
])

add_heading(doc,'四 面向工商银行的场景化适配',1)
add_heading(doc,'4.1 智能研究引擎',2)
add_para(doc,'弈金可以作为研究应用层，把公司公告、行业资料、公开信息和研究任务组织起来，形成“问题—证据—分析—复核—报告”的研究底稿。它的价值在于把搜索和模型回答变成可以查看、复核和复用的工作流，不与工行已有底层模型体系形成替代关系。')
add_heading(doc,'4.2 科技金融与企业客户分析',2)
add_para(doc,'围绕科技型企业或产业链企业，可以将行业景气、企业基本面、公开舆情、知识产权/资质材料和经营风险组织为一份辅助尽调报告。系统只提供结构化分析和待复核事项，授信、审批和客户决策仍由有权限的业务人员完成。')
add_heading(doc,'4.3 风险与金融安全',2)
add_para(doc,'金融场景必须把“安全”写进产品协议。弈金拟从数据、模型、应用和人工操作四个层面建立约束：文件按用户和机构隔离；敏感信息不进入普通日志；关键数字要求来源；模型判断和原始数据分层；报告导出前必须人工复核；系统不执行交易、不自动发布金融结论。')
add_heading(doc,'4.4 与工银智涌等公开方向的关系',2)
add_table(doc,['公开方向','弈金可承担的应用层角色','需要进一步验证的事项'],[
    ['多模型与金融级底座','通过 Provider 抽象承接模型能力，保持 Agent 编排层可替换。','模型接入标准、权限、部署边界和性能指标。'],
    ['企业知识工程','将研究资料、公开来源和任务结果按证据关系组织。','知识库范围、数据分级、更新机制和检索质量。'],
    ['智能体协同','以 Commander + Specialist 表达任务分工和依赖。','与既有智能体平台的协议、编排和审计接口。'],
    ['全链路安全防护','提供引用、任务、人工复核和导出留痕。','安全测试、渗透测试、模型红队和生产审批。'],
])

add_heading(doc,'五 技术架构与创新点',1)
add_heading(doc,'5.1 总体架构',2)
add_para(doc,'系统由 Web 前端、API 层、任务编排层、质量与证据层和基础设施组成。API 层负责认证、权限、文件、任务、Agent 和报告接口；任务编排层负责 Commander、DAG 调度、Specialist 执行和报告合成；质量与证据层负责文件解析、检索、引用核验、计算核验和报告审查；PostgreSQL、Redis、Worker 与对象存储提供持久化、队列和文件保存能力。')
add_heading(doc,'5.2 Commander 与 Specialist 协议',2)
add_table(doc,['组件','职责','输出'],[
    ['Commander','识别公司、研究目标和分析范围；选择合法 Agent；生成任务图。','任务 ID、Agent 名称、输入、输出、优先级、依赖关系。'],
    ['Specialist','读取任务和证据，按专业框架分析并自检。','内容、证据、数字核验、质量状态、待复核项。'],
    ['Search Provider','联网搜索、保存来源、去重并记录获取时间。','URL、标题、摘要、定位、抓取时间。'],
    ['Report Synthesizer','整合专家结果，标记冲突、失败、超时和跳过。','研究报告草稿、来源、风险提示和复核状态。'],
])
add_heading(doc,'5.3 三项创新',2)
add_bullet(doc,'工作流创新：将金融研究中的专业分工建模为 Commander + Specialist，而非一次性问答。')
add_bullet(doc,'证据链创新：把来源追溯、过程追溯和任务追溯作为统一协议，贯穿每个 Agent。')
add_bullet(doc,'落地创新：采用模型可替换、搜索独立、SaaS 与私有化并行的架构，降低金融机构试点门槛。')

add_heading(doc,'六 可行性、商业模式与竞争分析',1)
add_heading(doc,'6.1 技术可行性',2)
add_para(doc,'当前项目已具备 Web 前端、FastAPI 后端、Agent 注册与编排、SSE 流式输出、搜索引用持久化、用户登录、Docker Compose 部署等 MVP 基础。后续技术重点不是继续堆叠 Agent 数量，而是补齐结构化 DAG、真实公开数据源、质量评测和权限测试。')
add_heading(doc,'6.2 商业模式',2)
add_table(doc,['模式','服务对象','价值来源','前置条件'],[
    ['机构私有化','中小银行、券商、基金和研究团队。','按部署、席位和服务收费；数据不出机构边界。','通过安全评估、部署适配和项目验收。'],
    ['团队 SaaS','小型投研团队和高校金融实验室。','按席位、任务量或高级功能订阅。','公开数据和模型成本可控。'],
    ['能力服务','金融科技集成商或业务部门。','提供编排、证据链和报告工作流 API。','统一 Provider、Search 和审计接口。'],
])
add_heading(doc,'6.3 竞争差异',2)
add_table(doc,['对比维度','通用聊天工具','低代码 Agent 平台','弈金'],[
    ['金融方法论','需要用户自行组织。','需要自行配置。','内置行业、基本面、估值、风险等专业角色。'],
    ['任务协作','通常是一问一答。','灵活但搭建成本高。','Commander 自动拆解，Specialist 按依赖执行。'],
    ['证据与审计','来源和过程不稳定。','取决于用户配置。','统一保存引用、任务状态、质量状态和待复核项。'],
    ['部署与模型','受平台和模型限制。','需要工程维护。','Provider 解耦，支持 SaaS 与私有化方向。'],
])

add_heading(doc,'七 实施路径、评测与预算',1)
add_heading(doc,'7.1 实施路线',2)
add_table(doc,['阶段','主要任务','验收指标'],[
    ['阶段一：MVP 闭环','文件、任务、Worker、Agent 编排、引用、报告和人工复核。','可上传、可执行、可恢复、可回链、可导出。'],
    ['阶段二：可靠性增强','结构化 DAG、取消/重试、权限隔离、质量检查和错误提示。','关键路径测试；失败可见且可恢复。'],
    ['阶段三：数据与评测','接入公开结构化数据；建立 10–20 家 A 股公司评测集。','关键数字可追溯率目标 100%；引用覆盖率目标不低于 90%。'],
    ['阶段四：机构试点','私有化部署、安全评估、业务流程适配和人工反馈闭环。','通过部署验收和业务人员可用性测试。'],
])
add_heading(doc,'7.2 评测体系',2)
add_bullet(doc,'任务拆解：评测 Commander 是否选对 Agent、输入输出是否完整、依赖关系是否合理。')
add_bullet(doc,'事实与数字：检查关键数字、单位、期间、口径与来源定位。')
add_bullet(doc,'引用质量：检查来源可访问性、引用覆盖率、重复来源和待核验标记。')
add_bullet(doc,'工程可靠性：检查任务成功率、重试、取消、断线恢复、权限隔离和导出一致性。')
add_bullet(doc,'人工价值：记录耗时、失败原因和人工修改点，不用“模型说得像”作为唯一评价标准。')
add_heading(doc,'7.3 预算原则',2)
add_para(doc,'校赛阶段以低成本验证闭环为目标，主要成本包括模型调用、云服务器、域名与公开数据源。预算采用“按量计费、优先免费公开数据源、必要时私有化”的原则，不在项目早期预设昂贵的数据采购或复杂运维系统。进入机构试点后，再根据并发量、数据等级、安全测试和部署环境形成正式预算。')

add_heading(doc,'八 风险分析与合规边界',1)
add_table(doc,['风险','可能影响','控制措施'],[
    ['模型幻觉与引用失真','研究结论错误，降低用户信任。','强制来源、数字核验和待核验状态；人工复核后才能形成正式成果。'],
    ['数据口径不一致','不同来源期间、单位或统计口径不同。','保存原始数据、单位、期间、来源和计算公式；冲突时显式标记。'],
    ['敏感数据泄露','用户材料跨组织访问或进入普通日志。','用户/机构权限隔离；对象存储抽象；密钥只由环境变量注入；敏感信息不写普通日志。'],
    ['模型或搜索服务不可用','任务失败、超时或结果不完整。','后台 Worker、状态持久化、有限重试、失败原因记录和可恢复任务。'],
    ['金融误用','用户将草稿当作投资建议或自动决策。','明确风险提示；不自动交易；不生成个性化买卖指令；保留 Human-in-the-Loop。'],
    ['工行场景误读','将公开方向误写成合作或授权。','所有内容使用“适配建议/候选场景”表述，正式接入须另行授权与评估。'],
])

add_heading(doc,'九 项目效益与答辩演示',1)
add_heading(doc,'9.1 预期效益',2)
add_table(doc,['受益方','预期价值'],[
    ['研究人员','减少资料搜集、初步整理和报告排版时间，将注意力转向判断和复核。'],
    ['金融机构','沉淀统一的任务、证据和质量协议，为研究、尽调和风险场景提供可复用的应用层。'],
    ['工商银行场景','可与智能研究、科技金融、知识工程和智能体协同方向形成应用层对接假设。'],
    ['社会与教育','推动大学生把 AI 能力放入真实金融工作流，形成可评测、可解释、可复核的实践样本。'],
])
add_heading(doc,'9.2 现场 Demo 脚本',2)
add_table(doc,['时刻','演示动作','评委看到的价值'],[
    ['1','输入“研究某 A 股公司”，提交任务。','系统自动理解问题并生成任务图。'],
    ['2','展示 Commander 分派行业、基本面、消息面和风险任务。','多 Agent 协作不是概念，而是可见的执行过程。'],
    ['3','展示 SSE 状态和中间输出。','长任务在后台运行，用户能看到进度和失败状态。'],
    ['4','点击引用链接或 PDF 页码定位。','结论可以回到来源，避免只看模型措辞。'],
    ['5','打开报告草稿和复核项，导出 PDF。','系统保留人工边界，报告可带走、可回溯。'],
])

add_heading(doc,'十 结论',1)
add_para(doc,'弈金选择一个清晰而克制的切入点：面向金融机构的深度研究工作台。它不以 Agent 数量或模型参数为卖点，而以任务组织、证据链、质量审查、人工复核和可部署性回答金融 AI 的真实落地问题。')
add_para(doc,'在工商银行等大型金融机构已经建设金融大模型、知识工程和智能体协同能力的背景下，弈金的机会不在于重复建设底层模型，而在于提出一个可用于研究与尽调场景的应用层工作流方案。项目将以 MVP 闭环为基础，继续通过公开材料、标准评测集和可用性测试验证价值，逐步形成面向金融机构的可靠 AI 工作台。')

add_heading(doc,'附录 参考资料与事实边界',1)
add_source(doc,'赛事官方：第十七届“工行杯”全国大学生金融科技创新大赛介绍','https://www.gonghangbei.com/dsjs.html')
add_source(doc,'赛事官方：赛事指引与常见问题','https://www.gonghangbei.com/sszy.html')
add_source(doc,'赛事官方：往期作品','https://www.gonghangbei.com/index/Lists/index.html?id=38')
add_source(doc,'中国工商银行：2025 年年度报告','https://v.icbc.com.cn/userfiles/resources/icbcltd/download/2026/2025AnnualReportA.pdf')
add_source(doc,'中国工商银行：2025 年经营发展向新向优','https://www.icbc.com.cn/page/1211022925399216128.html')
add_para(doc,'事实边界说明：赛事规则、工行公开信息和行业判断可能随时间更新；本文仅引用公开资料，未包含工商银行非公开数据、内部系统信息或未获授权的合作表述。项目中的市场规模、效率和成本等指标如无可复现测量，将作为目标或假设，不作为已被真实用户验证的结论。', first=False)

OUT.parent.mkdir(parents=True, exist_ok=True); doc.save(OUT); print(OUT)
