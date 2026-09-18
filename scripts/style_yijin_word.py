from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/Users/laurence/Documents/金融agent')
DOCX = ROOT / 'docs/competition/策划书-弈金-数字金融版.docx'
MASCOT = ROOT / 'docs/competition/assets/yijin-mascot.png'
BG1 = ROOT / 'docs/competition/assets/body-background-icbc-ai.png'
BG2 = ROOT / 'docs/competition/assets/body-background-research-evidence.png'
BG3 = ROOT / 'docs/competition/assets/body-background-ai-orchestration.png'

NAVY = '18243F'
RED = 'C21230'
GOLD = 'B1853B'
GRAY = '667085'


def set_run(run, size=11, bold=False, color='202938', font='Songti SC'):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color='D9D9D9', size='6'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    run._r.addnext(fld)


def make_image_paragraph(doc, image, width=5.5, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(str(image), width=Inches(width))
    element = p._p
    doc._body._body.remove(element)
    return element


doc = Document(str(DOCX))
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.35)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.55)
section.right_margin = Cm(2.55)

# 全局正文样式
for style_name in ('Normal', 'Body Text'):
    style = doc.styles[style_name]
    style.font.name = 'Songti SC'
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Songti SC')
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor.from_string('202938')
    style.paragraph_format.line_spacing = 1.35
    style.paragraph_format.space_after = Pt(6)

for level, size, color in ((1, 17, NAVY), (2, 13, NAVY), (3, 11, GRAY)):
    style = doc.styles[f'Heading {level}']
    style.font.name = 'Songti SC'
    style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Songti SC')
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True

# 表格统一为深蓝表头、浅色交替行
for table in doc.tables:
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_border(cell)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.1
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    set_run(run, size=9, color='202938')
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_run(run, size=9, bold=True, color='FFFFFF')
            elif row_index % 2 == 0:
                set_cell_shading(cell, 'F3F5F8')

# 页眉页脚
header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.text = ''
set_run(header.add_run('弈金 · 工行杯数字金融项目策划书'), size=8.5, color=GRAY, font='Arial')
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.text = ''
set_run(footer.add_run('让金融研究更快、更稳、更可核验  ·  '), size=8.5, color=GRAY, font='Arial')
add_page_number(footer)

# 在最前面插入新的可编辑封面
body = doc._body._body
cover = []
for text, size, color, bold, space_before in [
    ('第十七届“工行杯”全国大学生金融科技创新大赛', 16, RED, True, 80),
    ('弈金', 38, NAVY, True, 28),
    ('面向银行投研场景的可追溯上市公司智能研究工作台', 17, RED, False, 4),
    ('数字金融项目策划书', 15, NAVY, False, 28),
    ('2026 年 9 月', 10, GRAY, False, 75),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    set_run(p.add_run(text), size=size, color=color, bold=bold)
    cover.append(p._p)
mascot_p = doc.add_paragraph()
mascot_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
mascot_p.paragraph_format.space_before = Pt(12)
mascot_p.add_run().add_picture(str(MASCOT), width=Inches(1.55))
cover.append(mascot_p._p)
break_p = doc.add_paragraph()
break_p.add_run().add_break()
cover.append(break_p._p)
for element in cover:
    body.remove(element)
for element in reversed(cover):
    body.insert(0, element)

# 在关键章节插入三张可编辑的章节视觉图，而非把整页 PDF 栅格化。
insertions = [
    ('二 行业背景与项目必要性', BG1),
    ('三 金融研究业务链条与痛点分析', BG2),
    ('七 技术架构与创新点', BG3),
]
for marker, image in insertions:
    for paragraph in list(doc.paragraphs):
        if paragraph.text.strip().startswith(marker):
            element = make_image_paragraph(doc, image, width=5.25)
            paragraph._p.addprevious(element)
            break

DOCX.unlink(missing_ok=True)
doc.save(DOCX)
print(DOCX)
